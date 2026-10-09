"""Parse the official 2024 Republic of Altai municipal validation tables.

The shipment workbook is the primary source because its annual sheet contains
11-digit OKTMO codes.  The tourism workbook has names only and is joined solely
through an exact, conservatively normalised match to the names in the shipment
workbook.  Suppressed and missing cells remain missing; they are never zeroed.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


REGION_NAME = "Республика Алтай"
ANNUAL_SHEET = "4 кв. 2024"
ANNUAL_PERIOD_LABEL = "Январь-декабрь 2024"
SHIPMENT_TOTAL_CODE = "101.АГ"
OKTMO_RE = re.compile(r"^\d{11}$")
SECTOR_RE = re.compile(r"^[A-S]$")
SUPPRESSED_RE = re.compile(r"^\s*(?:\.{3}|…)")

SHIPMENT_URL = (
    "https://22.rosstat.gov.ru/storage/mediabank/"
    "Отгрузка по МО (5005)_2024(1).xlsx"
)
TOURISM_URL = (
    "https://22.rosstat.gov.ru/storage/mediabank/"
    "t_oktmo_reg_БД ПМО_999_публ(1)(1).xlsx"
)

TOURISM_COLUMNS = {
    2: ("accommodation_establishments", "Число КСР", "units"),
    3: ("rooms", "Число номеров", "units"),
    4: ("places", "Число мест", "units"),
    5: ("overnight_stays", "Число ночевок - всего", "overnight_stays"),
    6: (
        "persons_accommodated",
        "Численность размещенных лиц - всего",
        "persons",
    ),
}

TOURISM_AGGREGATES = {
    "муниципальные образования республики алтай",
    "муниципальные районы республики алтай",
    "городские округа республики алтай/",
}


def normalise_name(value: object) -> str:
    """Normalise representation without removing meaningful words or suffixes."""
    text = unicodedata.normalize("NFKC", str(value)).replace("\xa0", " ")
    return " ".join(text.split()).casefold()


def normalise_oktmo(value: object) -> str:
    """Remove the display hyphens/spaces and validate an 11-digit OKTMO."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        raise ValueError("Missing OKTMO")
    text = re.sub(r"[-\s]", "", str(value).strip())
    if not OKTMO_RE.fullmatch(text):
        raise ValueError(f"Invalid 11-digit OKTMO: {value!r}")
    return text


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def parse_numeric_cell(value: object) -> tuple[float | None, str]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "missing"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value), "observed"
    if SUPPRESSED_RE.match(str(value)):
        return None, "suppressed"
    raise ValueError(f"Unexpected numeric cell: {value!r}")


def load_panel_crosswalk(panel_path: Path) -> pd.DataFrame:
    panel = pd.read_csv(
        panel_path,
        dtype={"territory_id": "string", "oktmo": "string", "region_name": "string"},
        usecols=["territory_id", "municipal_district_name", "region_name", "oktmo"],
    )
    panel = panel.loc[panel["region_name"] == REGION_NAME].copy()
    if panel.empty:
        raise ValueError(f"No {REGION_NAME!r} rows in {panel_path}")
    panel["oktmo"] = panel["oktmo"].map(normalise_oktmo)

    conflicts = panel.groupby("oktmo", dropna=False).agg(
        territory_ids=("territory_id", "nunique"),
        names=("municipal_district_name", "nunique"),
    )
    if (conflicts[["territory_ids", "names"]] != 1).any().any():
        raise ValueError(f"Non-unique panel mapping:\n{conflicts}")

    result = panel[
        ["territory_id", "oktmo", "municipal_district_name"]
    ].drop_duplicates()
    if len(result) != 11 or result["oktmo"].nunique() != 11:
        raise ValueError(f"Expected 11 unique Altai OKTMO codes in panel, got {len(result)}")
    return result.sort_values("oktmo").reset_index(drop=True)


def parse_shipments(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if ANNUAL_SHEET not in workbook.sheetnames:
            raise ValueError(f"Missing annual sheet {ANNUAL_SHEET!r}")
        sheet = workbook[ANNUAL_SHEET]
        if str(sheet["C4"].value).strip() != ANNUAL_PERIOD_LABEL:
            raise ValueError(f"Unexpected annual period: {sheet['C4'].value!r}")

        records: list[dict[str, object]] = []
        municipalities: list[dict[str, str]] = []
        current: dict[str, str] | None = None
        for name, code, raw_value in sheet.iter_rows(
            min_row=6, max_col=3, values_only=True
        ):
            code_text = "" if code is None else str(code).strip()
            if OKTMO_RE.fullmatch(code_text):
                current = {
                    "oktmo": normalise_oktmo(code_text),
                    "shipment_municipality_name": str(name).strip(),
                }
                municipalities.append(current.copy())
                continue
            if code_text != SHIPMENT_TOTAL_CODE and not SECTOR_RE.fullmatch(code_text):
                continue
            if current is None:
                raise ValueError(f"Indicator {code_text!r} appears before a municipality")

            value, status = parse_numeric_cell(raw_value)
            records.append(
                {
                    **current,
                    "year": 2024,
                    "period": "2024-01-01/2024-12-31",
                    "sector_code": "TOTAL" if code_text == SHIPMENT_TOTAL_CODE else code_text,
                    "sector_label": str(name).strip(),
                    "is_total": code_text == SHIPMENT_TOTAL_CODE,
                    "shipment_thousand_rubles": value,
                    "value_status": status,
                    "unit": "thousand_rubles",
                    "population_scope": "organizations_excluding_small_businesses",
                }
            )
    finally:
        workbook.close()

    municipality_frame = pd.DataFrame(municipalities).drop_duplicates()
    if len(municipality_frame) != 11 or municipality_frame["oktmo"].nunique() != 11:
        raise ValueError(
            f"Expected 11 unique OKTMO codes in shipment workbook, got {len(municipality_frame)}"
        )
    if municipality_frame["shipment_municipality_name"].map(normalise_name).duplicated().any():
        raise ValueError("Shipment municipality names are not unique after normalisation")

    shipment_frame = pd.DataFrame(records)
    if shipment_frame.duplicated(["oktmo", "sector_code"]).any():
        raise ValueError("Duplicate municipality-sector rows in annual shipment sheet")
    total_counts = shipment_frame.loc[shipment_frame["is_total"]].groupby("oktmo").size()
    if len(total_counts) != 11 or not total_counts.eq(1).all():
        raise ValueError("Every municipality must have exactly one shipment total")
    return shipment_frame, municipality_frame


def join_shipments_to_panel(
    shipments: pd.DataFrame, municipalities: pd.DataFrame, panel: pd.DataFrame
) -> tuple[pd.DataFrame, pd.DataFrame]:
    source_codes = set(municipalities["oktmo"])
    panel_codes = set(panel["oktmo"])
    if source_codes != panel_codes:
        raise ValueError(
            "Shipment and panel OKTMO sets differ: "
            f"source_only={sorted(source_codes - panel_codes)}, "
            f"panel_only={sorted(panel_codes - source_codes)}"
        )

    bridge = panel.merge(municipalities, on="oktmo", how="inner", validate="one_to_one")
    joined = shipments.merge(
        bridge[["territory_id", "oktmo", "municipal_district_name"]],
        on="oktmo",
        how="left",
        validate="many_to_one",
    )
    if joined["territory_id"].isna().any():
        raise ValueError("Unmatched shipment OKTMO after panel join")
    return joined, bridge


def parse_tourism(path: Path, bridge: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    name_bridge = bridge[
        ["territory_id", "oktmo", "shipment_municipality_name"]
    ].copy()
    name_bridge["normalised_name"] = name_bridge["shipment_municipality_name"].map(
        normalise_name
    )
    if name_bridge["normalised_name"].duplicated().any():
        raise ValueError("Tourism bridge names are not unique")
    lookup = name_bridge.set_index("normalised_name").to_dict("index")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError("Expected one tourism worksheet")
        sheet = workbook[workbook.sheetnames[0]]
        title = str(sheet["A1"].value or "")
        if "за 2024 год" not in title:
            raise ValueError(f"Unexpected tourism title: {title!r}")

        records: list[dict[str, object]] = []
        matched_rows: list[dict[str, object]] = []
        unexpected_names: list[str] = []
        for row in sheet.iter_rows(min_row=6, max_col=6, values_only=True):
            source_name = row[0]
            if source_name is None:
                continue
            normalised = normalise_name(source_name)
            if normalised in TOURISM_AGGREGATES:
                continue
            target = lookup.get(normalised)
            if target is None:
                unexpected_names.append(str(source_name))
                continue
            matched_rows.append(
                {
                    **target,
                    "tourism_municipality_name": str(source_name).strip(),
                    "normalised_name": normalised,
                }
            )
            for column, (indicator, label, unit) in TOURISM_COLUMNS.items():
                value, status = parse_numeric_cell(row[column - 1])
                records.append(
                    {
                        "territory_id": target["territory_id"],
                        "oktmo": target["oktmo"],
                        "tourism_municipality_name": str(source_name).strip(),
                        "year": 2024,
                        "indicator": indicator,
                        "indicator_label": label,
                        "value": value,
                        "value_status": status,
                        "unit": unit,
                    }
                )
    finally:
        workbook.close()

    if unexpected_names:
        raise ValueError(
            "Unmatched non-aggregate tourism names: " + ", ".join(unexpected_names)
        )
    matched = pd.DataFrame(matched_rows)
    if len(matched) != 11 or matched["oktmo"].nunique() != 11:
        raise ValueError(f"Expected 11 exact tourism name matches, got {len(matched)}")
    tourism = pd.DataFrame(records)
    if tourism.duplicated(["oktmo", "indicator"]).any():
        raise ValueError("Duplicate municipality-indicator rows in tourism workbook")
    return tourism, matched.sort_values("oktmo").reset_index(drop=True)


def build_outputs(
    shipment_path: Path,
    tourism_path: Path | None,
    panel_path: Path,
    output_dir: Path,
) -> dict[str, object]:
    panel = load_panel_crosswalk(panel_path)
    shipments, municipalities = parse_shipments(shipment_path)
    shipments, bridge = join_shipments_to_panel(shipments, municipalities, panel)

    output_dir.mkdir(parents=True, exist_ok=True)
    shipment_output = output_dir / "shipments_by_sector.csv"
    bridge_output = output_dir / "municipality_bridge.csv"
    shipments.to_csv(shipment_output, index=False, encoding="utf-8")

    tourism = None
    tourism_matches = None
    if tourism_path is not None and tourism_path.exists():
        tourism, tourism_matches = parse_tourism(tourism_path, bridge)
        tourism.to_csv(output_dir / "tourism.csv", index=False, encoding="utf-8")
        bridge = bridge.merge(
            tourism_matches[["oktmo", "tourism_municipality_name", "normalised_name"]],
            on="oktmo",
            how="left",
            validate="one_to_one",
        )
    bridge.to_csv(bridge_output, index=False, encoding="utf-8")

    metadata: dict[str, object] = {
        "scope": {
            "region": REGION_NAME,
            "year": 2024,
            "municipality_level": "first_level",
            "municipalities": 11,
            "national_generalisation": False,
        },
        "shipments": {
            "source_url": SHIPMENT_URL,
            "source_file": shipment_path.name,
            "sha256": sha256_file(shipment_path),
            "rows": int(len(shipments)),
            "observed_rows": int((shipments["value_status"] == "observed").sum()),
            "suppressed_rows": int((shipments["value_status"] == "suppressed").sum()),
            "missing_rows": int((shipments["value_status"] == "missing").sum()),
            "unit": "thousand_rubles",
            "population_scope": "organizations_excluding_small_businesses",
            "interpretation": (
                "Observed shipment by OKVED2 section; it is not employment, value added, "
                "or the full municipal economy. Suppressed and absent cells are not zero."
            ),
        },
        "panel": {
            "source_file": panel_path.as_posix(),
            "sha256": sha256_file(panel_path),
            "join": "11-digit OKTMO after removing display hyphens and spaces",
            "matched_municipalities": int(bridge["oktmo"].nunique()),
        },
        "tourism": None,
    }
    if tourism is not None and tourism_path is not None:
        metadata["tourism"] = {
            "source_url": TOURISM_URL,
            "source_file": tourism_path.name,
            "sha256": sha256_file(tourism_path),
            "rows": int(len(tourism)),
            "matched_municipalities": int(tourism["oktmo"].nunique()),
            "join": (
                "Exact NFKC/casefold/whitespace-normalised name match to the shipment "
                "workbook; no suffix removal or fuzzy matching."
            ),
            "interpretation": (
                "Counts of collective accommodation establishments and their activity; "
                "not all tourism activity."
            ),
        }
    (output_dir / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return metadata


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--shipments",
        type=Path,
        default=root
        / "artifacts/sources/external-20260923/altai_republic_shipments_2024.xlsx",
    )
    parser.add_argument(
        "--tourism",
        type=Path,
        default=root
        / "artifacts/sources/external-20260923/altai_republic_tourism_2024.xlsx",
    )
    parser.add_argument(
        "--panel", type=Path, default=root / "data/processed/panel.csv"
    )
    parser.add_argument(
        "--output", type=Path, default=root / "reports/external/altai-2024"
    )
    args = parser.parse_args()
    metadata = build_outputs(args.shipments, args.tourism, args.panel, args.output)
    print(json.dumps(metadata, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
